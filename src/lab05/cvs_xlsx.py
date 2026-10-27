import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from src.lab04.io_txt_csv import ensure_parent_dir

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX.
    Использует openpyxl.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """
    src = Path(csv_path)
    dst = Path(xlsx_path)

    if not src.exists():
        raise FileNotFoundError("Указанный CSV файл не найден")
    if src.suffix.lower() != ".csv":
        raise ValueError(f"Неверный тип файла: ожидается .csv, получено '{src.suffix}'")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError(f"Неверный тип файла: ожидается .xlsx, получено '{dst.suffix}'")

    with src.open('r', newline='', encoding='utf-8') as f:
        reader = list(csv.reader(f))

    if not reader:
        raise ValueError("Файл CSV пустой")

    widths = {}
    for row in reader:
        for index, cell in enumerate(row):
            if len(cell) > widths.get(index, 0):
                widths[index] = len(cell)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in reader:
        ws.append(row)

    for index, width in widths.items():
        letter = get_column_letter(index + 1)
        ws.column_dimensions[letter].width = max(8, width)

    ensure_parent_dir(dst)
    wb.save(dst)


if __name__ == '__main__':
    csv_to_xlsx(r"data/lab05/samples/people.csv", r"data/lab05/out/people.xlsx")
    csv_to_xlsx(r"data/lab05/samples/cities.csv", r"data/lab05/out/cities.xlsx")